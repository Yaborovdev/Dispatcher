import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell import MergedCell
from openpyxl.utils import get_column_letter

from services.templates import TEMPLATES, TemplateSpec

log = logging.getLogger(__name__)

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"

MONTHS_RU = [
    "", "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]


def _set_value(cell, value) -> None:
    cell.value = value
    # openpyxl превращает строку, начинающуюся с "=", в формулу.
    # Данные из БД должны попадать в файл как текст, а не как формула.
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def safe_write(ws, cell_address: str, value):
    """
    Безопасно записывает значение в ячейку.
    Если ячейка является частью объединённого диапазона (MergedCell),
    запись перенаправляется в левую верхнюю ячейку этого диапазона.
    """
    cell = ws[cell_address]

    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            anchor = (
                f"{get_column_letter(merged_range.min_col)}"
                f"{merged_range.min_row}"
            )
            _set_value(ws[anchor], value)
            return

    # MergedCell, но диапазон не найден — пропускаем во избежание ошибки
    if isinstance(cell, MergedCell):
        log.warning("%s — объединённая ячейка, но диапазон не найден. Запись пропущена.", cell_address)
        return

    _set_value(cell, value)


def build_values(waybill) -> dict[str, object]:
    """Собирает значения полей путевого листа; какие из них попадут в файл, решает шаблон."""
    values: dict[str, object] = {"number": waybill.number}

    d = waybill.date_issue
    if d:
        values.update(
            issue_day=d.day,
            issue_month=MONTHS_RU[d.month],
            issue_year=d.year,
            issue_year_short=str(d.year)[-2:],
        )

    start = waybill.date_start or waybill.date_issue
    if start:
        values.update(start_day=start.day, start_month=MONTHS_RU[start.month], start_year=start.year)

    end = waybill.date_end
    if end:
        values.update(end_day=end.day, end_month=MONTHS_RU[end.month], end_year=end.year)

    if waybill.vehicle:
        values.update(vehicle_model=waybill.vehicle.brand_model, vehicle_number=waybill.vehicle.gov_number)

    if waybill.driver:
        values.update(
            driver_name=waybill.driver.full_name,
            driver_license=waybill.driver.driver_license,
            driver_snils=waybill.driver.snils,
        )

    if waybill.trailer:
        values.update(trailer_model=waybill.trailer.brand_model, trailer_number=waybill.trailer.reg_number)

    if waybill.customer:
        values["customer"] = waybill.customer.name
    if waybill.route:
        values["route"] = waybill.route

    return values


def _get_spec(template_type: str) -> TemplateSpec:
    try:
        return TEMPLATES[template_type]
    except KeyError:
        raise ValueError(
            f"Неизвестный тип путевого листа: «{template_type}». "
            f"Измените тип в путевом листе или добавьте его в services/templates.py."
        ) from None


def fill_waybill_template(waybill, output_path: str | Path) -> Path:
    """
    Заполняет Excel-шаблон данными путевого листа.
    waybill - объект модели Waybill (со всеми связями driver, vehicle и т.д.)
    """
    spec = _get_spec(waybill.template_type)
    template_path = TEMPLATES_DIR / spec.file

    if not template_path.exists():
        raise FileNotFoundError(
            f"Для типа «{waybill.template_type}» нет шаблона: ожидается файл {template_path}"
        )
    if not spec.cells:
        raise ValueError(
            f"Для типа «{waybill.template_type}» не заданы ячейки шаблона (services/templates.py)."
        )

    wb = load_workbook(template_path)
    if spec.sheet not in wb.sheetnames:
        raise ValueError(
            f"В шаблоне {spec.file} нет листа «{spec.sheet}» (есть: {', '.join(wb.sheetnames)})."
        )
    ws = wb[spec.sheet]

    values = build_values(waybill)
    for field, address in spec.cells.items():
        if field in values:
            safe_write(ws, address, values[field])

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
